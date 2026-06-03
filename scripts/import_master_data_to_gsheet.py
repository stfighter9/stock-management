#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import OrderedDict
from pathlib import Path
from typing import Iterable

from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "shopee_odoo_mapping_result.xlsx"
DEFAULT_KEY_FILE = ROOT / "smart-sandbox-424309-n5-c615a9ae9610.json"

TARGET_SHEETS = OrderedDict(
    [
        ("Products", ["odoo_product_key", "display_name", "quantity_on_hand", "unit"]),
        ("Shopee_Catalog", ["sku", "product_name", "variation_name", "stock"]),
        ("Shopee_Mapping", ["shopee_sku", "mapping_type", "odoo_product_key", "mix_group_id", "conversion_qty", "active", "note"]),
        ("Shopee_Mapping_Components", ["shopee_sku", "component_no", "odoo_product_key", "component_qty", "active", "note"]),
        ("Mix_Color_Options", ["mix_group_id", "odoo_product_key", "display_name", "active"]),
    ]
)


def normalize(value) -> str:
    return "" if value is None else str(value).strip()


def number_or_blank(value):
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def iter_rows_as_dicts(ws) -> Iterable[dict[str, object]]:
    rows = ws.iter_rows(values_only=True)
    headers = [normalize(v) for v in next(rows)]
    for row in rows:
        if not any(cell not in (None, "") for cell in row):
            continue
        yield {headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))}


def parse_component_mapping(raw_value: object) -> list[str]:
    value = normalize(raw_value)
    if not value:
        return []
    items = []
    for line in value.splitlines():
        text = line.strip()
        if not text:
            continue
        items.append(text.split("|", 1)[0].strip())
    return items


def build_datasets(workbook_path: Path) -> dict[str, list[list[object]]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    mapping_rows = list(iter_rows_as_dicts(wb["Mapping_Result"]))
    odoo_rows = list(iter_rows_as_dicts(wb["Odoo_Products"]))

    products = [
        [
            normalize(row.get("odoo_product_key")),
            normalize(row.get("display_name")),
            number_or_blank(row.get("quantity_on_hand")),
            normalize(row.get("unit")),
        ]
        for row in odoo_rows
        if normalize(row.get("odoo_product_key"))
    ]

    shopee_catalog = [
        [
            normalize(row.get("SKU")),
            normalize(row.get("Product Name")),
            normalize(row.get("Variation Name")),
            "",
        ]
        for row in mapping_rows
    ]

    shopee_mapping: list[list[object]] = []
    mix_options_seen: OrderedDict[tuple[str, str], tuple[str, str, str, str]] = OrderedDict()

    for row in mapping_rows:
        sku = normalize(row.get("SKU"))
        mapping_type = normalize(row.get("suggested_mapping_type"))
        status = normalize(row.get("mapping_status"))
        if not sku or mapping_type not in {"FIXED_SKU", "MIX_COLOR"} or not status.startswith("AUTO_"):
            continue

        mapped_odoo_product_key = normalize(row.get("mapped_odoo_product_key"))
        mix_group_id = normalize(row.get("mix_group_id"))
        conversion_qty = number_or_blank(row.get("conversion_qty"))
        note = normalize(row.get("mapping_note"))

        shopee_mapping.append(
            [
                sku,
                mapping_type,
                mapped_odoo_product_key if mapping_type == "FIXED_SKU" else "",
                mix_group_id if mapping_type == "MIX_COLOR" else "",
                conversion_qty,
                "TRUE",
                note,
            ]
        )

        if mapping_type == "MIX_COLOR":
            for odoo_product_key in parse_component_mapping(row.get("component_mapping")):
                key = (mix_group_id, odoo_product_key)
                if key not in mix_options_seen:
                    mix_options_seen[key] = (mix_group_id, odoo_product_key, odoo_product_key, "TRUE")

    mix_color_options = [list(values) for values in mix_options_seen.values()]
    shopee_mapping_components: list[list[object]] = []

    return {
        "Products": products,
        "Shopee_Catalog": shopee_catalog,
        "Shopee_Mapping": shopee_mapping,
        "Shopee_Mapping_Components": shopee_mapping_components,
        "Mix_Color_Options": mix_color_options,
    }


def ensure_sheets(service, spreadsheet_id: str) -> None:
    spreadsheet = service.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
    existing = {sheet["properties"]["title"] for sheet in spreadsheet.get("sheets", [])}
    requests = []
    for name, headers in TARGET_SHEETS.items():
        if name in existing:
            continue
        requests.append(
            {
                "addSheet": {
                    "properties": {
                        "title": name,
                        "gridProperties": {
                            "rowCount": 1000,
                            "columnCount": max(len(headers), 8),
                            "frozenRowCount": 1,
                        },
                    }
                }
            }
        )
    if requests:
        service.spreadsheets().batchUpdate(spreadsheetId=spreadsheet_id, body={"requests": requests}).execute()


def clear_target_sheets(service, spreadsheet_id: str) -> None:
    service.spreadsheets().values().batchClear(
        spreadsheetId=spreadsheet_id,
        body={"ranges": [f"{name}!A2:Z" for name in TARGET_SHEETS]},
    ).execute()


def write_datasets(service, spreadsheet_id: str, datasets: dict[str, list[list[object]]]) -> None:
    data = []
    for name, headers in TARGET_SHEETS.items():
        rows = datasets[name]
        values = [headers, *rows]
        data.append({"range": f"{name}!A1", "values": values})

    service.spreadsheets().values().batchUpdate(
        spreadsheetId=spreadsheet_id,
        body={"valueInputOption": "USER_ENTERED", "data": data},
    ).execute()


def build_service(key_file: Path):
    credentials = Credentials.from_service_account_file(str(key_file), scopes=["https://www.googleapis.com/auth/spreadsheets"])
    return build("sheets", "v4", credentials=credentials)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import master stock data into Google Sheets.")
    parser.add_argument("--spreadsheet-id", required=True, help="Target Google Spreadsheet ID")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK), help="Path to shopee_odoo_mapping_result.xlsx")
    parser.add_argument("--key-file", default=str(DEFAULT_KEY_FILE), help="Path to Google service account JSON")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook).resolve()
    key_file = Path(args.key_file).resolve()

    datasets = build_datasets(workbook_path)
    service = build_service(key_file)
    ensure_sheets(service, args.spreadsheet_id)
    clear_target_sheets(service, args.spreadsheet_id)
    write_datasets(service, args.spreadsheet_id, datasets)

    for name, rows in datasets.items():
        print(f"{name}: {len(rows)} rows")


if __name__ == "__main__":
    main()
