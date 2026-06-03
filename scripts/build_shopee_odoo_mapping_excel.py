#!/usr/bin/env python3
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SHOPEE_FILE = ROOT / "shopee-product.xlsx"
ODOO_FILE = ROOT / "odoo-product.xlsx"
OUT_FILE = ROOT / "shopee_odoo_mapping_result.xlsx"

NS = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

MODEL_PATTERNS = [
    r"FO-[A-Z0-9]+",
    r"C-[A-Z0-9]+(?:/[A-Z0-9]+)?",
    r"CAL-[A-Z0-9]+(?:/[A-Z0-9]+)?",
    r"WACO-[A-Z0-9]+",
    r"POSCO-[A-Z0-9]+",
]
BAD_MODEL_CODES = {
    "ANTI-LEAKAGE",
    "ANTI-SLIP",
    "DUAL-NIB",
    "LONG-LASTING",
    "LOW-ODOR",
    "MULTI-SURFACE",
    "MULTI-SURFACES",
    "NON-TOXIC",
    "QUICK-DRY",
    "WOOD-FREE",
}
COLORS = {
    "black": ["black", "blk"],
    "blue": ["blue", "blu"],
    "red": ["red"],
    "green": ["green"],
    "yellow": ["yellow"],
    "orange": ["orange"],
    "pink": ["pink"],
    "purple": ["purple"],
    "violet": ["violet"],
    "brown": ["brown"],
    "magenta": ["magenta"],
    "turquoise": ["turquoise"],
    "grey": ["grey", "gray"],
    "clear": ["clear", "transparent"],
    "white": ["white"],
    "light blue": ["light blue", "l.blue", "l blue"],
    "light green": ["light green", "l.green", "l green"],
}


@dataclass
class OdooProduct:
    display_name: str
    quantity_on_hand: float
    unit: str
    model_codes: list[str]
    colors: set[str]
    size_tokens: set[str]


def normalize_text(value: Any) -> str:
    return str(value or "").strip()


def normalize_number(value: Any) -> float:
    try:
        return float(str(value or "0").replace(",", "").strip())
    except ValueError:
        return 0.0


def col_index(cell_ref: str) -> int:
    letters = "".join(ch for ch in cell_ref if ch.isalpha())
    index = 0
    for letter in letters:
        index = index * 26 + ord(letter.upper()) - 64
    return index - 1


def read_xlsx_sheet1(path: Path) -> list[list[str]]:
    with ZipFile(path) as archive:
        shared_strings: list[str] = []
        try:
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall("a:si", NS):
                shared_strings.append("".join(t.text or "" for t in item.findall(".//a:t", NS)))
        except KeyError:
            pass

        sheet_root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        rows: list[list[str]] = []
        for row in sheet_root.findall(".//a:sheetData/a:row", NS):
            values: list[str] = []
            for cell in row.findall("a:c", NS):
                index = col_index(cell.attrib["r"])
                while len(values) <= index:
                    values.append("")

                cell_type = cell.attrib.get("t")
                value_node = cell.find("a:v", NS)
                if cell_type == "s":
                    values[index] = shared_strings[int(value_node.text)] if value_node is not None and value_node.text else ""
                elif cell_type == "inlineStr":
                    values[index] = "".join(t.text or "" for t in cell.findall(".//a:t", NS))
                else:
                    values[index] = value_node.text if value_node is not None and value_node.text else ""
            rows.append(values)
        return rows


def model_code_aliases(code: str) -> set[str]:
    code = code.upper().strip()
    aliases = {code}
    if "/" in code:
        aliases.add(code.split("/", 1)[0])
    return aliases


def extract_model_codes(text: str) -> list[str]:
    upper = text.upper()
    codes: list[str] = []
    for pattern in MODEL_PATTERNS:
        codes.extend(re.findall(pattern, upper))
    output = []
    seen = set()
    for code in codes:
        for alias in model_code_aliases(code):
            if alias not in BAD_MODEL_CODES and alias not in seen:
                seen.add(alias)
                output.append(alias)
    return output


def extract_colors(text: str) -> set[str]:
    normalized = f" {text.lower().replace('-', ' ')} "
    dotted = f" {text.lower()} "
    found: set[str] = set()
    for color, aliases in COLORS.items():
        for alias in aliases:
            haystack = dotted if "." in alias else normalized
            if re.search(r"\b" + re.escape(alias) + r"\b", haystack):
                found.add(color)
                if color == "light blue":
                    found.add("blue")
                if color == "light green":
                    found.add("green")
    return found


def extract_size_tokens(text: str) -> set[str]:
    lower = text.lower()
    tokens = set(re.findall(r"\b(?:0\.[0-9]+|1\.0|[0-9]+x[0-9]+m|[0-9]+mm)\b", lower))
    return {token.replace("mm", "") for token in tokens}


def parse_pack_qty(variation_name: str, product_name: str) -> int:
    text = variation_name or product_name
    lower = text.lower()
    patterns = [
        r"bag\s*(\d+)\s*pcs?",
        r"box\s*(\d+)\s*pcs?",
        r"box\s*(\d+)pcs?",
        r"pack\s*(\d+)\s*pcs?",
        r"(\d+)\s*pcs?",
        r"set\s*(\d+)\s*(?:colors?|pcs?)",
        r"(\d+)\s*colors?\s*set",
    ]
    for pattern in patterns:
        match = re.search(pattern, lower)
        if match:
            return int(match.group(1))
    return 1


def parse_pack_count(variation_name: str) -> int | None:
    match = re.search(r"\b(\d+)\s*packs?\b", variation_name.lower())
    return int(match.group(1)) if match else None


def conversion_qty_for_odoo_unit(odoo_product: OdooProduct, product_name: str, variation_name: str) -> int:
    if "pack" in odoo_product.unit.lower():
        return parse_pack_count(variation_name) or parse_pack_count(product_name) or 1
    return parse_pack_qty(variation_name, product_name)


def is_assorted_or_random(variation_name: str, product_name: str) -> bool:
    target = variation_name if variation_name else product_name
    return bool(re.search(r"\b(assorted|random|mix|mixed)\b", target, re.IGNORECASE))


def detect_tip_size(product_name: str, variation_name: str) -> str:
    variation_sizes = re.findall(r"\b(0\.[0-9]+|1\.0)\s*mm\b", variation_name.lower())
    if variation_sizes:
        return variation_sizes[-1]

    product_sizes = re.findall(r"\b(0\.[0-9]+|1\.0)\s*mm\b", product_name.lower())
    if product_sizes:
        return product_sizes[-1]
    return ""


def normalize_display_lookup(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def set_fixed_mapping(
    result: dict[str, Any],
    odoo_product: OdooProduct,
    conversion_qty: int,
    note: str,
    status: str = "AUTO_MAPPED_FO_GELB08",
) -> dict[str, Any]:
    result["suggested_mapping_type"] = "FIXED_SKU"
    result["mapped_odoo_product_key"] = odoo_product.display_name
    result["mapped_odoo_display_name"] = odoo_product.display_name
    result["odoo_qty_on_hand"] = odoo_product.quantity_on_hand
    result["odoo_unit"] = odoo_product.unit
    result["conversion_qty"] = conversion_qty
    result["mapping_status"] = status
    result["mapping_confidence"] = "HIGH"
    result["mapping_note"] = note
    return result


def set_combo_mapping(
    result: dict[str, Any],
    component_rows: list[tuple[OdooProduct, int]],
    note: str,
    status: str,
) -> dict[str, Any]:
    result["suggested_mapping_type"] = "COMBO_SKU"
    result["conversion_qty"] = sum(component_qty for _, component_qty in component_rows)
    result["component_mapping"] = "\n".join(
        f"{item.display_name} | component_qty={component_qty}" for item, component_qty in component_rows
    )
    result["component_odoo_qty_on_hand"] = "\n".join(
        f"{item.display_name} | odoo_qty_on_hand={item.quantity_on_hand:g} {item.unit}"
        for item, _ in component_rows
    )
    result["mapping_status"] = status
    result["mapping_confidence"] = "HIGH"
    result["mapping_note"] = note
    return result


def set_mix_mapping(
    result: dict[str, Any],
    component_rows: list[OdooProduct],
    conversion_qty: int,
    mix_group_id: str,
    note: str,
    status: str,
) -> dict[str, Any]:
    result["suggested_mapping_type"] = "MIX_COLOR"
    result["conversion_qty"] = conversion_qty
    result["mix_group_id"] = mix_group_id
    result["component_mapping"] = "\n".join(item.display_name for item in component_rows)
    result["component_odoo_qty_on_hand"] = "\n".join(
        f"{item.display_name} | odoo_qty_on_hand={item.quantity_on_hand:g} {item.unit}"
        for item in component_rows
    )
    result["mapping_status"] = status
    result["mapping_confidence"] = "HIGH"
    result["mapping_note"] = note
    return result


def map_assorted_combo_if_exact(
    result: dict[str, Any],
    product_name: str,
    variation_name: str,
    component_names: list[str],
    odoo_by_display: dict[str, OdooProduct],
    mix_group_id: str,
    note: str,
    status: str,
    unsupported_status: str,
    unsupported_note: str,
    missing_status: str,
) -> dict[str, Any]:
    pack_qty = parse_pack_qty(variation_name, product_name)
    if pack_qty % len(component_names) != 0:
        result["mapping_status"] = unsupported_status
        result["mapping_note"] = unsupported_note
        return result

    missing = [name for name in component_names if normalize_display_lookup(name) not in odoo_by_display]
    if missing:
        result["mapping_status"] = missing_status
        result["mapping_note"] = f"Expected combo components not found: {', '.join(missing)}"
        return result

    components = [odoo_by_display[normalize_display_lookup(name)] for name in component_names]
    return set_mix_mapping(result, components, pack_qty, mix_group_id, note, status)


def apply_fo_gelb08_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GELB08" not in " ".join(extract_model_codes(text)):
        return None

    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GELB08 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    lower_text = text.lower()
    lower_variation = variation_name.lower()
    tip_size = detect_tip_size(product_name, variation_name)
    colors = extract_colors(variation_name) or extract_colors(product_name)

    def exists(display_name: str) -> bool:
        return normalize_display_lookup(display_name) in odoo_by_display

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any] | None:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(result, product, conversion_qty, note)
        result["mapping_status"] = "REVIEW_NO_FO_GELB08_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GELB08 Odoo variant not found: {display_name}"
        return result

    if "pack 6pcs" in lower_text or "assorted 6 colors" in lower_variation:
        pack_count = parse_pack_count(variation_name)
        if pack_count:
            conversion_qty = pack_count
        elif "box 12pcs" in lower_variation:
            conversion_qty = 2
        else:
            conversion_qty = 1
        return map_if_exists(
            "FO-GELB08/PH FLEXSTICK 0.7 multi colors pack 6s/100",
            conversion_qty,
            "FO-GELB08 assorted 6-color pack maps to Odoo pack-of-6 product.",
        )

    if "pack 3 colors" in lower_text:
        size = tip_size or "0.7"
        if size == "0.5":
            display = "FO-GELB08/PH FLEXSTICK 0.5 Blk-Blu-Red pack 3s/200"
        else:
            display = "FO-GELB08/PH FLEXSTICK 0.7 Blk-Blu-Red pack 3s/200"
        return map_if_exists(display, 1, "FO-GELB08 fixed 3-color pack maps to Odoo pack-of-3 product.")

    conversion_qty = parse_pack_qty(variation_name, product_name)
    color_to_display = {
        ("0.5", "black"): "FO-GELB08 FLEXSTICK 0.5 Black box 12s/1200",
        ("0.5", "blue"): "FO-GELB08 FLEXSTICK 0.5 Blue box 12s/1200",
        ("0.5", "red"): "FO-GELB08 FLEXSTICK 0.5 Red box 12s/1200",
        ("0.7", "black"): "FO-GELB08 FLEXSTICK 0.7 Black box 12s/1200",
        ("0.7", "blue"): "FO-GELB08 FLEXSTICK 0.7 Blue box 12s/1200",
        ("0.7", "red"): "FO-GELB08 FLEXSTICK 0.7 Red box 12s/1200",
        ("0.7", "green"): "FO-GELB08 FLEXSTICK 0.7 Green box 12s/1200",
        ("0.7", "light blue"): "FO-GELB08/PH FLEXSTICK 0.7 L.Blue box 12s/1200",
        ("0.7", "light green"): "FO-GELB08/PH FLEXSTICK 0.7 L.Green box 12s/1200",
        ("0.7", "yellow"): "FO-GELB08/PH FLEXSTICK 0.7 Gold box 12s/1200",
        ("0.7", "orange"): "FO-GELB08/PH FLEXSTICK 0.7 Orange box 12s/1200",
        ("0.7", "pink"): "FO-GELB08/PH FLEXSTICK 0.7 Pink box 12s/1200",
        ("0.7", "purple"): "FO-GELB08/PH FLEXSTICK 0.7 Purple box 12s/1200",
        ("1.0", "black"): "FO-GELB08 FLEXSTICK 1.0 Black box 12s/1200",
        ("1.0", "blue"): "FO-GELB08 FLEXSTICK 1.0 Blue box 12s/1200",
        ("1.0", "red"): "FO-GELB08 FLEXSTICK 1.0 Red box 12s/1200",
    }

    color_priority = ["light blue", "light green", "black", "blue", "red", "green", "yellow", "orange", "pink", "purple"]
    for color in color_priority:
        if color in colors:
            display = color_to_display.get((tip_size, color))
            if display:
                return map_if_exists(
                    display,
                    conversion_qty,
                    f"FO-GELB08 single-color {tip_size}mm maps by color and pack quantity.",
                )

    result["mapping_status"] = "REVIEW_FO_GELB08_UNRESOLVED"
    result["mapping_note"] = "FO-GELB08 row could not be mapped to a known Odoo variant by color/size/pack rules."
    return result


def apply_fo_gelb036_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GELB036" not in " ".join(extract_model_codes(text)):
        return None

    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GELB036 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def exists(display_name: str) -> bool:
        return normalize_display_lookup(display_name) in odoo_by_display

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any] | None:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(
                result,
                product,
                conversion_qty,
                note,
                status="AUTO_MAPPED_FO_GELB036",
            )
        result["mapping_status"] = "REVIEW_NO_FO_GELB036_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GELB036 Odoo variant not found: {display_name}"
        return result

    tip_size = detect_tip_size(product_name, variation_name)
    colors = extract_colors(variation_name) or extract_colors(product_name)
    conversion_qty = parse_pack_qty(variation_name, product_name)
    is_assorted = is_assorted_or_random(variation_name, product_name)

    color_to_display = {
        ("0.5", "black"): "FO-GELB036/PH FLEXTOK 0.5 Black box 12s/1200",
        ("0.5", "blue"): "FO-GELB036/PH FLEXTOK 0.5 Blue box 12s/1200",
        ("0.5", "red"): "FO-GELB036/PH FLEXTOK 0.5 Red box 12s/1200",
        ("0.7", "black"): "FO-GELB036/PH FLEXTOK 0.7 Black box 12s/1200",
        ("0.7", "blue"): "FO-GELB036/PH FLEXTOK 0.7 Blue box 12s/1200",
        ("0.7", "red"): "FO-GELB036/PH FLEXTOK 0.7 Red box 12s/1200",
        ("1.0", "black"): "FO-GELB036/PH FLEXTOK 1.0 Black box 12s/1200",
        ("1.0", "blue"): "FO-GELB036/PH FLEXTOK 1.0 Blue box 12s/1200",
        ("1.0", "red"): "FO-GELB036/PH FLEXTOK 1.0 Red box 12s/1200",
    }

    if is_assorted and tip_size == "0.5":
        return map_if_exists(
            "FO-GELB036/PH FLEXTOK 0.5 Blk Assorted box 12s/1200",
            conversion_qty,
            "FO-GELB036 0.5 assorted maps to the dedicated Odoo assorted item.",
        )

    if is_assorted and tip_size == "0.7":
        pack_qty = parse_pack_qty(variation_name, product_name)
        if pack_qty not in {3, 6, 12}:
            result["mapping_status"] = "REVIEW_FO_GELB036_ASSORTED_QTY"
            result["mapping_note"] = "FO-GELB036 0.7 assorted row has unsupported pack quantity."
            return result

        each_qty = pack_qty // 3
        component_names = [
            "FO-GELB036/PH FLEXTOK 0.7 Black box 12s/1200",
            "FO-GELB036/PH FLEXTOK 0.7 Blue box 12s/1200",
            "FO-GELB036/PH FLEXTOK 0.7 Red box 12s/1200",
        ]
        missing = [name for name in component_names if not exists(name)]
        if missing:
            result["mapping_status"] = "REVIEW_NO_FO_GELB036_ODOO_VARIANT"
            result["mapping_note"] = f"Expected FO-GELB036 combo components not found: {', '.join(missing)}"
            return result
        components = [odoo_by_display[normalize_display_lookup(name)] for name in component_names]
        return set_mix_mapping(
            result,
            components,
            pack_qty,
            "MIX_FO_GELB036_3COLORS",
            "FO-GELB036 0.7 assorted maps to a Black/Blue/Red color pool because there is no dedicated assorted Odoo SKU.",
            "AUTO_MIX_FO_GELB036",
        )

    color_priority = ["black", "blue", "red"]
    for color in color_priority:
        if color in colors:
            display = color_to_display.get((tip_size, color))
            if display:
                return map_if_exists(
                    display,
                    conversion_qty,
                    f"FO-GELB036 single-color {tip_size}mm maps by color and pack quantity.",
                )

    result["mapping_status"] = "REVIEW_FO_GELB036_UNRESOLVED"
    result["mapping_note"] = "FO-GELB036 row could not be mapped to a known Odoo variant by color/size/pack rules."
    return result


def apply_fo_gel04_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GEL04" not in extract_model_codes(text):
        return None

    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GEL04 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def exists(display_name: str) -> bool:
        return normalize_display_lookup(display_name) in odoo_by_display

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any] | None:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(
                result,
                product,
                conversion_qty,
                note,
                status="AUTO_MAPPED_FO_GEL04",
            )
        result["mapping_status"] = "REVIEW_NO_FO_GEL04_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GEL04 Odoo variant not found: {display_name}"
        return result

    conversion_qty = parse_pack_qty(variation_name, product_name)
    colors = extract_colors(variation_name) or extract_colors(product_name)
    color_to_display = {
        "black": "FO-GEL04 SUNBEAM 0.5 Black box 12s/600",
        "blue": "FO-GEL04 SUNBEAM 0.5 Blue box 12s/600",
        "green": "FO-GEL04 SUNBEAM 0.5 Green box 12s/600",
        "purple": "FO-GEL04 SUNBEAM 0.5 Purple box 12s/600",
        "red": "FO-GEL04 SUNBEAM 0.5 Red box 12s/600",
    }

    if is_assorted_or_random(variation_name, product_name):
        result["mapping_status"] = "REVIEW_FO_GEL04_ASSORTED_NO_EXACT_MATCH"
        result["mapping_note"] = "FO-GEL04 assorted 5-color row is left blank because Shopee includes Pink but Odoo has no exact Pink SUNBEAM variant."
        return result

    if "pink" in colors:
        result["mapping_status"] = "REVIEW_FO_GEL04_NO_EXACT_ODOO_COLOR"
        result["mapping_note"] = "FO-GEL04 Pink row is left blank because Odoo has no exact Pink SUNBEAM variant."
        return result

    for color in ["black", "blue", "green", "purple", "red"]:
        if color in colors:
            return map_if_exists(
                color_to_display[color],
                conversion_qty,
                "FO-GEL04 single-color row maps to the exact SUNBEAM normal color variant.",
            )

    result["mapping_status"] = "REVIEW_FO_GEL04_UNRESOLVED"
    result["mapping_note"] = "FO-GEL04 row could not be mapped to a known Odoo variant by color/pack rules."
    return result


def apply_fo_gele002_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GELE002" not in extract_model_codes(text):
        return None

    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GELE002 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any] | None:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(
                result,
                product,
                conversion_qty,
                note,
                status="AUTO_MAPPED_FO_GELE002",
            )
        result["mapping_status"] = "REVIEW_NO_FO_GELE002_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GELE002 Odoo variant not found: {display_name}"
        return result

    colors = extract_colors(variation_name) or extract_colors(product_name)
    conversion_qty = parse_pack_qty(variation_name, product_name)
    color_to_display = {
        "black": "FO-GELE002 FLEXCORRECT 0.5 Black box 12s/600",
        "blue": "FO-GELE002 FLEXCORRECT 0.5 Blue box 12s/600",
    }

    if "violet" in colors:
        result["mapping_status"] = "REVIEW_FO_GELE002_NO_EXACT_ODOO_COLOR"
        result["mapping_note"] = "FO-GELE002 Violet row is left blank because Odoo has no exact Violet FLEXCORRECT variant."
        return result

    for color in ["black", "blue"]:
        if color in colors:
            return map_if_exists(
                color_to_display[color],
                conversion_qty,
                "FO-GELE002 single-color row maps to the exact FLEXCORRECT box 12s/600 color variant.",
            )

    result["mapping_status"] = "REVIEW_FO_GELE002_UNRESOLVED"
    result["mapping_note"] = "FO-GELE002 row could not be mapped to a known Odoo variant by exact color/pack rules."
    return result


def apply_fo_gelb06_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GELB06" not in extract_model_codes(text):
        return None

    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GELB06 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def exists(display_name: str) -> bool:
        return normalize_display_lookup(display_name) in odoo_by_display

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any] | None:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(
                result,
                product,
                conversion_qty,
                note,
                status="AUTO_MAPPED_FO_GELB06",
            )
        result["mapping_status"] = "REVIEW_NO_FO_GELB06_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GELB06 Odoo variant not found: {display_name}"
        return result

    conversion_qty = parse_pack_qty(variation_name, product_name)
    lower_text = text.lower()
    colors = extract_colors(variation_name) or extract_colors(product_name)

    if "pastel barrel" in lower_text:
        return map_if_exists(
            "FO-GELB06 PSLIDE PASTEL 0.5 Blk Assorted box 12s/600",
            conversion_qty,
            "FO-GELB06 Pastel Barrel row maps to the exact PSLIDE PASTEL black-ink assorted-barrel Odoo variant.",
        )

    color_to_display = {
        "black": "FO-GELB06 PSLIDE 0.5 Black box 12s/600",
        "blue": "FO-GELB06 PSLIDE 0.5 Blue box 12s/600",
        "red": "FO-GELB06 PSLIDE 0.5 Red box 12s/600",
    }

    if is_assorted_or_random(variation_name, product_name):
        pack_qty = parse_pack_qty(variation_name, product_name)
        if pack_qty not in {3, 6, 12}:
            result["mapping_status"] = "REVIEW_FO_GELB06_ASSORTED_QTY"
            result["mapping_note"] = "FO-GELB06 assorted 3-color row has unsupported pack quantity."
            return result

        component_names = [
            "FO-GELB06 PSLIDE 0.5 Black box 12s/600",
            "FO-GELB06 PSLIDE 0.5 Blue box 12s/600",
            "FO-GELB06 PSLIDE 0.5 Red box 12s/600",
        ]
        missing = [name for name in component_names if not exists(name)]
        if missing:
            result["mapping_status"] = "REVIEW_NO_FO_GELB06_ODOO_VARIANT"
            result["mapping_note"] = f"Expected FO-GELB06 combo components not found: {', '.join(missing)}"
            return result
        components = [odoo_by_display[normalize_display_lookup(name)] for name in component_names]
        return set_mix_mapping(
            result,
            components,
            pack_qty,
            "MIX_FO_GELB06_3COLORS",
            "FO-GELB06 assorted 3-color row maps to a Black/Blue/Red color pool.",
            "AUTO_MIX_FO_GELB06",
        )

    for color in ["black", "blue", "red"]:
        if color in colors:
            return map_if_exists(
                color_to_display[color],
                conversion_qty,
                "FO-GELB06 single-color row maps to the exact PSLIDE normal color variant.",
            )

    result["mapping_status"] = "REVIEW_FO_GELB06_UNRESOLVED"
    result["mapping_note"] = "FO-GELB06 row could not be mapped to a known Odoo variant by exact color/pack rules."
    return result


def apply_fo_gel021_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GEL021" not in extract_model_codes(text):
        return None
    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GEL021 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any]:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(result, product, conversion_qty, note, status="AUTO_MAPPED_FO_GEL021")
        result["mapping_status"] = "REVIEW_NO_FO_GEL021_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GEL021 Odoo variant not found: {display_name}"
        return result

    conversion_qty = parse_pack_qty(variation_name, product_name)
    colors = extract_colors(variation_name) or extract_colors(product_name)
    lower_name = product_name.lower()

    if "blister" in lower_name:
        return map_if_exists(
            "FO-GEL021 G-MASTER 0.5 Black blister 1/96",
            1,
            "FO-GEL021 blister row maps to the exact black blister Odoo variant.",
        )

    if "refill" in lower_name and "gel pen" not in lower_name.replace("ink refill gel pen", ""):
        return map_if_exists(
            "Refill FO-GEL021/PH G-MASTER 0.5 Black box 12s/1728",
            conversion_qty,
            "FO-GEL021 refill row maps to the exact black refill Odoo variant.",
        )

    if is_assorted_or_random(variation_name, product_name):
        return map_assorted_combo_if_exact(
            result,
            product_name,
            variation_name,
            [
                "FO-GEL021 G-MASTER 0.5 Black box 12s/600",
                "FO-GEL021 G-MASTER 0.5 Blue box 12s/600",
                "FO-GEL021 G-MASTER 0.5 Red box 12s/600",
            ],
            odoo_by_display,
            "MIX_FO_GEL021_3COLORS",
            "FO-GEL021 assorted 3-color row maps to a Black/Blue/Red color pool.",
            "AUTO_MIX_FO_GEL021",
            "REVIEW_FO_GEL021_ASSORTED_QTY",
            "FO-GEL021 assorted 3-color row has unsupported pack quantity.",
            "REVIEW_NO_FO_GEL021_ODOO_VARIANT",
        )

    color_to_display = {
        "black": "FO-GEL021 G-MASTER 0.5 Black box 12s/600",
        "blue": "FO-GEL021 G-MASTER 0.5 Blue box 12s/600",
        "red": "FO-GEL021 G-MASTER 0.5 Red box 12s/600",
    }
    for color in ["black", "blue", "red"]:
        if color in colors:
            return map_if_exists(
                color_to_display[color],
                conversion_qty,
                "FO-GEL021 single-color row maps to the exact G-MASTER pen color variant.",
            )

    result["mapping_status"] = "REVIEW_FO_GEL021_UNRESOLVED"
    result["mapping_note"] = "FO-GEL021 row could not be mapped by exact product-form/color rules."
    return result


def apply_fo_gelb019_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GELB019" not in extract_model_codes(text):
        return None
    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GELB019 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any]:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(result, product, conversion_qty, note, status="AUTO_MAPPED_FO_GELB019")
        result["mapping_status"] = "REVIEW_NO_FO_GELB019_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GELB019 Odoo variant not found: {display_name}"
        return result

    tip_size = detect_tip_size(product_name, variation_name)
    if "pack 3pcs" in product_name.lower():
        display = f"FO-GELB019/PH FLEXSTICK NEO {tip_size} Black pack 3s/200"
        return map_if_exists(display, 1, "FO-GELB019 pack row maps to the exact pack-of-3 black variant.")

    conversion_qty = parse_pack_qty(variation_name, product_name)
    display = f"FO-GELB019/PH FLEXSTICK NEO {tip_size} Black box 12s/1200"
    return map_if_exists(display, conversion_qty, "FO-GELB019 single-color row maps to the exact black box variant.")


def apply_fo_gel022_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GEL022" not in extract_model_codes(text):
        return None
    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GEL022 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any]:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(result, product, conversion_qty, note, status="AUTO_MAPPED_FO_GEL022")
        result["mapping_status"] = "REVIEW_NO_FO_GEL022_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GEL022 Odoo variant not found: {display_name}"
        return result

    colors = extract_colors(variation_name) or extract_colors(product_name)
    conversion_qty = parse_pack_qty(variation_name, product_name)
    if "violet" in colors:
        result["mapping_status"] = "REVIEW_FO_GEL022_NO_EXACT_ODOO_COLOR"
        result["mapping_note"] = "FO-GEL022 Violet row is left blank because Odoo has no exact Violet FLEXGEL variant."
        return result

    if is_assorted_or_random(variation_name, product_name):
        result["mapping_status"] = "REVIEW_FO_GEL022_ASSORTED_NO_EXACT_MATCH"
        result["mapping_note"] = "FO-GEL022 assorted 5-color row is left blank because Odoo lacks an exact fifth color component."
        return result

    color_to_display = {
        "black": "FO-GEL022 FLEXGEL 0.5 Black box 12s/600",
        "blue": "FO-GEL022 FLEXGEL 0.5 Blue box 12s/600",
        "green": "FO-GEL022 FLEXGEL 0.5 Green box 12s/600",
        "red": "FO-GEL022 FLEXGEL 0.5 Red box 12s/600",
    }
    for color in ["black", "blue", "green", "red"]:
        if color in colors:
            return map_if_exists(color_to_display[color], conversion_qty, "FO-GEL022 single-color row maps to the exact FLEXGEL color variant.")

    result["mapping_status"] = "REVIEW_FO_GEL022_UNRESOLVED"
    result["mapping_note"] = "FO-GEL022 row could not be mapped by exact color rules."
    return result


def apply_fo_gele003_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-GELE003" not in extract_model_codes(text):
        return None
    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-GELE003 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any]:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(result, product, conversion_qty, note, status="AUTO_MAPPED_FO_GELE003")
        result["mapping_status"] = "REVIEW_NO_FO_GELE003_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-GELE003 Odoo variant not found: {display_name}"
        return result

    colors = extract_colors(variation_name) or extract_colors(product_name)
    conversion_qty = parse_pack_qty(variation_name, product_name)
    if "violet" in colors:
        result["mapping_status"] = "REVIEW_FO_GELE003_NO_EXACT_ODOO_COLOR"
        result["mapping_note"] = "FO-GELE003 Violet row is left blank because Odoo has no exact Violet FLEXMAZZIC variant."
        return result

    color_to_display = {
        "black": "FO-GELE003 FLEXMAZZIC 0.5 Black box 12s/600",
        "blue": "FO-GELE003 FLEXMAZZIC 0.5 Blue box 12s/600",
    }
    for color in ["black", "blue"]:
        if color in colors:
            return map_if_exists(color_to_display[color], conversion_qty, "FO-GELE003 single-color row maps to the exact FLEXMAZZIC color variant.")

    result["mapping_status"] = "REVIEW_FO_GELE003_UNRESOLVED"
    result["mapping_note"] = "FO-GELE003 row could not be mapped by exact color rules."
    return result


def apply_fo_042_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "FO-042" not in extract_model_codes(text):
        return None
    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "FO-042 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any]:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(result, product, conversion_qty, note, status="AUTO_MAPPED_FO_042")
        result["mapping_status"] = "REVIEW_NO_FO_042_ODOO_VARIANT"
        result["mapping_note"] = f"Expected FO-042 Odoo variant not found: {display_name}"
        return result

    conversion_qty = parse_pack_qty(variation_name, product_name)
    lower_variation = variation_name.lower()
    if "black - blue" in lower_variation:
        return map_if_exists("FO-042/PH TWIN CANDEE 0.6 Blk-Blu jar 50/1200", conversion_qty, "FO-042 pair row maps to the exact Black-Blue variant.")
    if "black - red" in lower_variation:
        return map_if_exists("FO-042/PH TWIN CANDEE 0.6 Blk-Red jar 50/1200", conversion_qty, "FO-042 pair row maps to the exact Black-Red variant.")

    result["mapping_status"] = "REVIEW_FO_042_UNRESOLVED"
    result["mapping_note"] = "FO-042 row could not be mapped by exact pair-color rules."
    return result


def apply_posco_03_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"

    if "POSCO-03" not in extract_model_codes(text):
        return None
    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "POSCO-03 row has empty Shopee SKU; mapping is not usable until SKU is filled."
        return result

    def map_if_exists(display_name: str, conversion_qty: int, note: str) -> dict[str, Any]:
        product = odoo_by_display.get(normalize_display_lookup(display_name))
        if product:
            return set_fixed_mapping(result, product, conversion_qty, note, status="AUTO_MAPPED_POSCO_03")
        result["mapping_status"] = "REVIEW_NO_POSCO_03_ODOO_VARIANT"
        result["mapping_note"] = f"Expected POSCO-03 Odoo variant not found: {display_name}"
        return result

    lower_variation = variation_name.lower()
    exact_sets = {
        "black - set 6pcs": "POSCO-03/XK Poster Color 15ml Black set 6s/60",
        "green - set 6pcs": "POSCO-03/XK Poster Color 15ml Green set 6s/60",
        "red - set 6pcs": "POSCO-03/XK Poster Color 15ml Red set 6s/60",
        "white - set 6pcs": "POSCO-03/XK Poster Color 15ml White set 6s/60",
        "yellow - set 6pcs": "POSCO-03/XK Poster Color 15ml Yellow set 6s/60",
    }
    for key, display_name in exact_sets.items():
        if key in lower_variation:
            return map_if_exists(display_name, 1, "POSCO-03 row maps to the exact set variant.")
    if "blue - set 6pcs" in lower_variation:
        result["mapping_status"] = "REVIEW_POSCO_03_NO_EXACT_ODOO_COLOR"
        result["mapping_note"] = "POSCO-03 Blue set row is left blank because Odoo only has Cobalt Blue and Sky Blue, not exact Blue."
        return result

    result["mapping_status"] = "REVIEW_POSCO_03_UNRESOLVED"
    result["mapping_note"] = "POSCO-03 row could not be mapped by exact color/unit rules."
    return result


def apply_generic_assorted_combo_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    text = f"{product_name} {variation_name}"
    model_codes = extract_model_codes(text)
    if not model_codes:
        return None
    code = model_codes[0]
    if not is_assorted_or_random(variation_name, product_name):
        return None

    if code == "FO-PM03":
        component_names = [
            "FO-PM03/XK Perma.Marker Bullet Black box 12s/600",
            "FO-PM03/XK Perma.Marker Blue box 12s/600",
            "FO-PM03/XK Perma.Marker Red box 12s/600",
        ]
        missing = [name for name in component_names if normalize_display_lookup(name) not in odoo_by_display]
        if missing:
            result["mapping_status"] = "REVIEW_NO_FO_PM03_ODOO_VARIANT"
            result["mapping_note"] = f"Expected FO-PM03 mix components not found: {', '.join(missing)}"
            return result
        components = [odoo_by_display[normalize_display_lookup(name)] for name in component_names]
        return set_mix_mapping(
            result,
            components,
            parse_pack_qty(variation_name, product_name),
            "MIX_FO_PM03_3COLORS",
            "FO-PM03 assorted row maps to a Black/Blue/Red color pool, not a fixed component bundle.",
            "AUTO_MIX_FO_PM03",
        )

    combo_map = {
        "FO-04": [
            "FO-04 TANGO 0.7 Black box 12s/600",
            "FO-04 TANGO 0.7 Blue box 12s/600",
            "FO-04 TANGO 0.7 Red box 12s/600",
        ],
        "FO-GEL069": [
            "FO-GEL069 FAVOREE 0.38 Black box 12s/600",
            "FO-GEL069 FAVOREE 0.38 Blue box 12s/600",
            "FO-GEL069 FAVOREE 0.38 Red box 12s/600",
        ],
        "FO-GELB014": [
            "FO-GELB014 LARIS 0.5 Black box 12s/600",
            "FO-GELB014 LARIS 0.5 Blue box 12s/600",
            "FO-GELB014 LARIS 0.5 Red box 12s/600",
        ],
        "FO-GELB017": [
            "FO-GELB017 MEGA 0.7 Black box 12s/600",
            "FO-GELB017 MEGA 0.7 Blue box 12s/600",
            "FO-GELB017 MEGA 0.7 Red box 12s/600",
        ],
        "FO-GELB047": [
            "FO-GELB047 FLEXSUN 0.5 Black box 12s/600",
            "FO-GELB047 FLEXSUN 0.5 Blue box 12s/600",
            "FO-GELB047 FLEXSUN 0.5 Red box 12s/600",
        ],
        "FO-GELB09": [
            "FO-GELB09/PH SUPER TRENDEE 0.5 Black box 12s/600",
            "FO-GELB09/PH SUPER TRENDEE 0.5 Blue box 12s/600",
            "FO-GELB09/PH SUPER TRENDEE 0.5 Red box 12s/600",
        ],
        "FO-PM01": [
            "FO-PM01/XK Perma.Marker Black box 12s/600",
            "FO-PM01/XK Perma.Marker Blue box 12s/600",
            "FO-PM01/XK Perma.Marker Red box 12s/600",
        ],
        "FO-PM014": [
            "FO-PM014/XK Perma.Marker Black box 12s/600",
            "FO-PM014/XK Perma.Marker Blue box 12s/600",
            "FO-PM014/XK Perma.Marker Red box 12s/600",
        ],
        "FO-PM06": [
            "FO-PM06/XK Dual Tip Perma.Marker Black box 12s/600",
            "FO-PM06/XK Dual Tip Perma.Marker Blue box 12s/600",
            "FO-PM06/XK Dual Tip Perma.Marker RED box 12s/600",
        ],
        "FO-PMI01": [
            "FO-PMI01 Perma.Marker Ink 25ml Black bottle 6s/144",
            "FO-PMI01 Perma.Marker Ink 25ml Blue bottle 6s/144",
            "FO-PMI01 Perma.Marker Ink 25ml Red bottle 6s/144",
        ],
        "FO-WB016": [
            "FO-WB016/XK Whiteboard Marker Black box 12s/600",
            "FO-WB016/XK Whiteboard Marker Blue box 12s/600",
            "FO-WB016/XK Whiteboard Marker Red box 12s/600",
        ],
        "FO-WB03": [
            "FO-WB03/XK Whiteboard Marker Black box 12s/600",
            "FO-WB03/XK Whiteboard Marker Blue box 12s/600",
            "FO-WB03/XK Whiteboard Marker Red box 12s/600",
        ],
        "FO-WBI01": [
            "FO-WBI01 Whiteboard Marker Ink 25ml Black bottle 6s/144",
            "FO-WBI01 Whiteboard Marker Ink 25ml Blue bottle 6s/144",
            "FO-WBI01 Whiteboard Marker Ink 25ml Red bottle 6s/144",
        ],
    }
    component_names = combo_map.get(code)
    if not component_names:
        return None

    return map_assorted_combo_if_exact(
        result,
        product_name,
        variation_name,
        component_names,
        odoo_by_display,
        f"MIX_{code.replace('-', '_').replace('/', '_')}_3COLORS",
        f"{code} assorted row maps to an exact color pool in Odoo.",
        f"AUTO_MIX_{code.replace('-', '_').replace('/', '_')}",
        f"REVIEW_{code.replace('-', '_').replace('/', '_')}_ASSORTED_QTY",
        f"{code} assorted row has unsupported pack quantity.",
        f"REVIEW_NO_{code.replace('-', '_').replace('/', '_')}_ODOO_VARIANT",
    )


def apply_single_candidate_fixed_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"
    model_codes = extract_model_codes(text)
    if not model_codes:
        return None
    code = model_codes[0]

    display_name = ""
    if code == "C-G01" and "random color" in variation_name.lower():
        display_name = "C-G01 Liquid Glue 30ml tray 12/240"
    elif code == "C-G02" and "random color" in variation_name.lower():
        display_name = "C-G02 Liquid Glue 60ml box 12s/240"
    elif code == "FO-HL009" and "pack of 5 colors" in product_name.lower():
        display_name = "FO-HL009 Dual Tip Highlighter 5 colors set 5s/600"
    elif code == "FO-HL016" and "pack of 5 colors" in product_name.lower():
        display_name = "FO-HL016 Dual Tip Highlighter 5 colors set 5s/600"
    if not display_name:
        return None

    product = odoo_by_display.get(normalize_display_lookup(display_name))
    if not product:
        result["mapping_status"] = f"REVIEW_NO_{code.replace('-', '_')}_ODOO_VARIANT"
        result["mapping_note"] = f"Expected {code} Odoo variant not found: {display_name}"
        return result

    conversion_qty = conversion_qty_for_odoo_unit(product, product_name, variation_name)
    return set_fixed_mapping(
        result,
        product,
        conversion_qty,
        f"{code} row maps to the exact single Odoo set/SKU variant.",
        status=f"AUTO_MAPPED_{code.replace('-', '_')}",
    )


def apply_specific_review_or_mix_mapping(row: dict[str, Any], result: dict[str, Any], odoo_by_display: dict[str, OdooProduct]) -> dict[str, Any] | None:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    text = f"{product_name} {variation_name}"
    model_codes = extract_model_codes(text)
    if not model_codes:
        return None
    code = model_codes[0]

    if code == "FO-CDT001":
        result["mapping_status"] = "REVIEW_FO_CDT001_NO_EXACT_COLOR"
        result["mapping_note"] = "FO-CDT001 row is left blank because Shopee does not specify Blue or Grey, while Odoo requires an exact color variant."
        return result

    if code == "FO-CDT002":
        result["mapping_status"] = "REVIEW_FO_CDT002_NO_EXACT_COLOR"
        result["mapping_note"] = "FO-CDT002 row is left blank because Shopee does not specify Blue or Grey, while Odoo requires an exact color variant."
        return result

    if code == "FO-SR006":
        colors = extract_colors(variation_name) or extract_colors(product_name)
        if "green" in colors:
            result["mapping_status"] = "REVIEW_FO_SR006_NO_EXACT_ODOO_COLOR"
            result["mapping_note"] = "FO-SR006 Green row is left blank because Odoo has only Blue, Pink, and Purple ruler variants."
            return result
        return None

    hl_mix_components = {
        "FO-HL03": [
            "FO-HL03 Highlighter Blue tub 12s/288",
            "FO-HL03 Highlighter Green tub 12s/288",
            "FO-HL03 Highlighter Orange tub 12s/288",
            "FO-HL03 Highlighter Pink tub 12s/288",
            "FO-HL03 Highlighter Yellow tub 12s/288",
        ],
        "FO-HL07": [
            "FO-HL07 Dual Tip Highlighter Blue tub 12s/600",
            "FO-HL07 Dual Tip Highlighter Green tub 12s/600",
            "FO-HL07 Dual Tip Highlighter Orange tub 12s/600",
            "FO-HL07 Dual Tip Highlighter Pink tub 12s/600",
            "FO-HL07 Dual Tip Highlighter Yellow tub 12s/600",
        ],
    }
    component_names = hl_mix_components.get(code)
    if component_names and is_assorted_or_random(variation_name, product_name):
        missing = [name for name in component_names if normalize_display_lookup(name) not in odoo_by_display]
        if missing:
            result["mapping_status"] = f"REVIEW_NO_{code.replace('-', '_')}_ODOO_VARIANT"
            result["mapping_note"] = f"Expected {code} mix components not found: {', '.join(missing)}"
            return result
        components = [odoo_by_display[normalize_display_lookup(name)] for name in component_names]
        return set_mix_mapping(
            result,
            components,
            parse_pack_qty(variation_name, product_name),
            f"MIX_{code.replace('-', '_')}_5COLORS",
            f"{code} assorted 5-color row maps to a 5-color stock pool with exact Odoo components.",
            f"AUTO_MIX_{code.replace('-', '_')}",
        )

    return None


def parse_shopee_rows(rows: list[list[str]]) -> tuple[list[str], list[dict[str, Any]]]:
    headers = rows[2]
    header_index = {header: index for index, header in enumerate(headers)}
    parsed = []
    for source_row, row in enumerate(rows[6:], start=7):
        if not any(normalize_text(value) for value in row):
            continue
        record = {
            header: normalize_text(row[index]) if index < len(row) else ""
            for header, index in header_index.items()
        }
        record["source_row"] = source_row
        parsed.append(record)
    return headers, parsed


def parse_odoo_products(rows: list[list[str]]) -> list[OdooProduct]:
    headers = rows[0]
    header_index = {header: index for index, header in enumerate(headers)}
    products = []
    for row in rows[1:]:
        if not any(normalize_text(value) for value in row):
            continue
        display_name = normalize_text(row[header_index["Display Name"]])
        model_codes = extract_model_codes(display_name)
        if not model_codes and display_name:
            model_codes = [display_name.split()[0].upper()]
        products.append(
            OdooProduct(
                display_name=display_name,
                quantity_on_hand=normalize_number(row[header_index["Quantity On Hand"]]),
                unit=normalize_text(row[header_index["Unit"]]),
                model_codes=model_codes,
                colors=extract_colors(display_name),
                size_tokens=extract_size_tokens(display_name),
            )
        )
    return products


def candidate_summary(candidates: list[OdooProduct]) -> str:
    return "\n".join(
        f"{item.display_name} | qty={item.quantity_on_hand:g} | unit={item.unit}"
        for item in candidates[:20]
    )


def match_row(row: dict[str, Any], odoo_by_model: dict[str, list[OdooProduct]], sku_counts: Counter[str]) -> dict[str, Any]:
    product_name = row.get("Product Name", "")
    variation_name = row.get("Variation Name", "")
    sku = row.get("SKU", "")
    text = f"{product_name} {variation_name}"
    model_codes = extract_model_codes(text)
    colors = extract_colors(variation_name) or extract_colors(product_name)
    size_tokens = extract_size_tokens(text)
    conversion_qty = parse_pack_qty(variation_name, product_name)

    candidates: list[OdooProduct] = []
    seen = set()
    for code in model_codes:
        for candidate in odoo_by_model.get(code, []):
            if candidate.display_name not in seen:
                seen.add(candidate.display_name)
                candidates.append(candidate)

    filtered = candidates
    if colors and candidates:
        color_filtered = [item for item in candidates if item.colors & colors]
        if color_filtered:
            filtered = color_filtered

    if size_tokens and len(filtered) > 1:
        size_filtered = [item for item in filtered if item.size_tokens & size_tokens]
        if size_filtered:
            filtered = size_filtered

    result = {
        "suggested_mapping_type": "",
        "mapped_odoo_product_key": "",
        "mapped_odoo_display_name": "",
        "odoo_qty_on_hand": "",
        "odoo_unit": "",
        "conversion_qty": "",
        "mix_group_id": "",
        "component_mapping": "",
        "component_odoo_qty_on_hand": "",
        "mapping_status": "",
        "mapping_confidence": "",
        "model_codes_detected": ", ".join(model_codes),
        "colors_detected": ", ".join(sorted(colors)),
        "candidate_count": len(candidates),
        "filtered_candidate_count": len(filtered),
        "odoo_candidates": candidate_summary(filtered or candidates),
        "mapping_note": "",
    }

    odoo_by_display = {
        normalize_display_lookup(item.display_name): item
        for candidates in odoo_by_model.values()
        for item in candidates
    }

    fo_gelb08_result = apply_fo_gelb08_mapping(row, result, odoo_by_display)
    if fo_gelb08_result is not None:
        return fo_gelb08_result

    fo_gelb036_result = apply_fo_gelb036_mapping(row, result, odoo_by_display)
    if fo_gelb036_result is not None:
        return fo_gelb036_result

    fo_gel04_result = apply_fo_gel04_mapping(row, result, odoo_by_display)
    if fo_gel04_result is not None:
        return fo_gel04_result

    fo_gele002_result = apply_fo_gele002_mapping(row, result, odoo_by_display)
    if fo_gele002_result is not None:
        return fo_gele002_result

    fo_gelb06_result = apply_fo_gelb06_mapping(row, result, odoo_by_display)
    if fo_gelb06_result is not None:
        return fo_gelb06_result

    fo_gel021_result = apply_fo_gel021_mapping(row, result, odoo_by_display)
    if fo_gel021_result is not None:
        return fo_gel021_result

    fo_gelb019_result = apply_fo_gelb019_mapping(row, result, odoo_by_display)
    if fo_gelb019_result is not None:
        return fo_gelb019_result

    fo_gel022_result = apply_fo_gel022_mapping(row, result, odoo_by_display)
    if fo_gel022_result is not None:
        return fo_gel022_result

    fo_gele003_result = apply_fo_gele003_mapping(row, result, odoo_by_display)
    if fo_gele003_result is not None:
        return fo_gele003_result

    fo_042_result = apply_fo_042_mapping(row, result, odoo_by_display)
    if fo_042_result is not None:
        return fo_042_result

    posco_03_result = apply_posco_03_mapping(row, result, odoo_by_display)
    if posco_03_result is not None:
        return posco_03_result

    generic_combo_result = apply_generic_assorted_combo_mapping(row, result, odoo_by_display)
    if generic_combo_result is not None:
        return generic_combo_result

    single_fixed_result = apply_single_candidate_fixed_mapping(row, result, odoo_by_display)
    if single_fixed_result is not None:
        return single_fixed_result

    specific_review_or_mix_result = apply_specific_review_or_mix_mapping(row, result, odoo_by_display)
    if specific_review_or_mix_result is not None:
        return specific_review_or_mix_result

    if not sku:
        result["mapping_status"] = "NO_SHOPEE_SKU"
        result["mapping_note"] = "Shopee row has empty SKU; PRD cannot map/process this row until SKU is filled."
        return result

    if sku_counts[sku] > 1:
        result["mapping_status"] = "DUPLICATE_SHOPEE_SKU_REVIEW"
        result["mapping_note"] = "Shopee SKU is duplicated in source export; review before using it as mapping key."
        return result

    if not model_codes:
        result["mapping_status"] = "NO_MODEL_CODE"
        result["mapping_note"] = "No product model code like FO-..., C-..., CAL-... was detected from Shopee name/variation."
        return result

    if not candidates:
        result["mapping_status"] = "NO_ODOO_CANDIDATE"
        result["mapping_note"] = "Detected model code but no Odoo Display Name contains the same code."
        return result

    if is_assorted_or_random(variation_name, product_name):
        result["mapping_status"] = "REVIEW_MIX_OR_COMBO"
        result["mix_group_id"] = f"MIX_{model_codes[0]}"
        result["mapping_note"] = "Assorted/random/mix row; decide MIX_COLOR if picker chooses actual Odoo items, or COMBO_SKU if components are fixed."
        return result

    if len(filtered) == 1:
        matched = filtered[0]
        result["suggested_mapping_type"] = "FIXED_SKU"
        result["mapped_odoo_product_key"] = matched.display_name
        result["mapped_odoo_display_name"] = matched.display_name
        result["odoo_qty_on_hand"] = matched.quantity_on_hand
        result["odoo_unit"] = matched.unit
        result["conversion_qty"] = conversion_qty_for_odoo_unit(matched, product_name, variation_name)
        result["mapping_status"] = "AUTO_MAPPED"
        result["mapping_confidence"] = "HIGH" if colors or size_tokens or len(candidates) == 1 else "MEDIUM"
        result["mapping_note"] = "Auto mapped by model code plus color/size filters."
        return result

    result["mapping_status"] = "REVIEW_AMBIGUOUS"
    result["mapping_note"] = "Multiple Odoo candidates remain after model/color/size filters; mapping left blank."
    return result


def autosize_columns(ws) -> None:
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = normalize_text(cell.value)
            max_length = max(max_length, min(len(value), 80))
        ws.column_dimensions[column_letter].width = max(max_length + 2, 12)


def write_workbook(shopee_headers: list[str], rows: list[dict[str, Any]], odoo_products: list[OdooProduct], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mapping_Result"

    mapping_headers = [
        "suggested_mapping_type",
        "mapped_odoo_product_key",
        "mapped_odoo_display_name",
        "odoo_qty_on_hand",
        "odoo_unit",
        "conversion_qty",
        "mix_group_id",
        "component_mapping",
        "component_odoo_qty_on_hand",
        "mapping_status",
        "mapping_confidence",
        "model_codes_detected",
        "colors_detected",
        "candidate_count",
        "filtered_candidate_count",
        "odoo_candidates",
        "mapping_note",
    ]
    headers = ["source_row"] + shopee_headers + mapping_headers
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize_columns(ws)

    fo_sheet = wb.create_sheet("FO_GELB08_Detail")
    fo_headers = [
        "source_row",
        "SKU",
        "Product Name",
        "Variation Name",
        "mapped_odoo_product_key",
        "odoo_qty_on_hand",
        "odoo_unit",
        "conversion_qty",
        "mapping_status",
        "mapping_note",
        "odoo_candidates",
    ]
    fo_sheet.append(fo_headers)
    for cell in fo_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        if "FO-GELB08" in row.get("model_codes_detected", ""):
            fo_sheet.append([row.get(header, "") for header in fo_headers])
    fo_sheet.freeze_panes = "A2"
    fo_sheet.auto_filter.ref = fo_sheet.dimensions
    autosize_columns(fo_sheet)

    fo036_sheet = wb.create_sheet("FO_GELB036_Detail")
    fo036_headers = [
        "source_row",
        "SKU",
        "Product Name",
        "Variation Name",
        "mapped_odoo_product_key",
        "odoo_qty_on_hand",
        "odoo_unit",
        "conversion_qty",
        "component_mapping",
        "component_odoo_qty_on_hand",
        "mapping_status",
        "mapping_note",
        "odoo_candidates",
    ]
    fo036_sheet.append(fo036_headers)
    for cell in fo036_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        if "FO-GELB036" in row.get("model_codes_detected", ""):
            fo036_sheet.append([row.get(header, "") for header in fo036_headers])
    fo036_sheet.freeze_panes = "A2"
    fo036_sheet.auto_filter.ref = fo036_sheet.dimensions
    autosize_columns(fo036_sheet)

    fo04_sheet = wb.create_sheet("FO_GEL04_Detail")
    fo04_headers = [
        "source_row",
        "SKU",
        "Product Name",
        "Variation Name",
        "mapped_odoo_product_key",
        "odoo_qty_on_hand",
        "odoo_unit",
        "conversion_qty",
        "mapping_status",
        "mapping_note",
        "odoo_candidates",
    ]
    fo04_sheet.append(fo04_headers)
    for cell in fo04_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        if row.get("model_codes_detected", "") == "FO-GEL04":
            fo04_sheet.append([row.get(header, "") for header in fo04_headers])
    fo04_sheet.freeze_panes = "A2"
    fo04_sheet.auto_filter.ref = fo04_sheet.dimensions
    autosize_columns(fo04_sheet)

    fo_gele002_sheet = wb.create_sheet("FO_GELE002_Detail")
    fo_gele002_headers = [
        "source_row",
        "SKU",
        "Product Name",
        "Variation Name",
        "mapped_odoo_product_key",
        "odoo_qty_on_hand",
        "odoo_unit",
        "conversion_qty",
        "mapping_status",
        "mapping_note",
        "odoo_candidates",
    ]
    fo_gele002_sheet.append(fo_gele002_headers)
    for cell in fo_gele002_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        if row.get("model_codes_detected", "") == "FO-GELE002":
            fo_gele002_sheet.append([row.get(header, "") for header in fo_gele002_headers])
    fo_gele002_sheet.freeze_panes = "A2"
    fo_gele002_sheet.auto_filter.ref = fo_gele002_sheet.dimensions
    autosize_columns(fo_gele002_sheet)

    fo06_sheet = wb.create_sheet("FO_GELB06_Detail")
    fo06_headers = [
        "source_row",
        "SKU",
        "Product Name",
        "Variation Name",
        "mapped_odoo_product_key",
        "odoo_qty_on_hand",
        "odoo_unit",
        "conversion_qty",
        "component_mapping",
        "component_odoo_qty_on_hand",
        "mapping_status",
        "mapping_note",
        "odoo_candidates",
    ]
    fo06_sheet.append(fo06_headers)
    for cell in fo06_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        if row.get("model_codes_detected", "") == "FO-GELB06":
            fo06_sheet.append([row.get(header, "") for header in fo06_headers])
    fo06_sheet.freeze_panes = "A2"
    fo06_sheet.auto_filter.ref = fo06_sheet.dimensions
    autosize_columns(fo06_sheet)

    summary = wb.create_sheet("Summary")
    status_counts = Counter(row["mapping_status"] for row in rows)
    summary.append(["metric", "value"])
    summary.append(["shopee_rows", len(rows)])
    summary.append(["odoo_rows", len(odoo_products)])
    for status, count in status_counts.most_common():
        summary.append([status, count])
    summary.append(["auto_mapped_non_empty", sum(1 for row in rows if row["mapped_odoo_product_key"])])
    autosize_columns(summary)

    odoo_sheet = wb.create_sheet("Odoo_Products")
    odoo_sheet.append(["odoo_product_key", "display_name", "quantity_on_hand", "unit", "model_codes", "colors", "size_tokens"])
    for cell in odoo_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for item in odoo_products:
        odoo_sheet.append(
            [
                item.display_name,
                item.display_name,
                item.quantity_on_hand,
                item.unit,
                ", ".join(item.model_codes),
                ", ".join(sorted(item.colors)),
                ", ".join(sorted(item.size_tokens)),
            ]
        )
    odoo_sheet.freeze_panes = "A2"
    odoo_sheet.auto_filter.ref = odoo_sheet.dimensions
    autosize_columns(odoo_sheet)

    wb.save(output_path)


def main() -> None:
    shopee_headers, shopee_rows = parse_shopee_rows(read_xlsx_sheet1(SHOPEE_FILE))
    odoo_products = parse_odoo_products(read_xlsx_sheet1(ODOO_FILE))

    odoo_by_model: dict[str, list[OdooProduct]] = defaultdict(list)
    for product in odoo_products:
        for code in product.model_codes:
            for alias in model_code_aliases(code):
                odoo_by_model[alias].append(product)

    sku_counts = Counter(row.get("SKU", "") for row in shopee_rows if row.get("SKU", ""))
    output_rows = []
    for row in shopee_rows:
        output_rows.append({**row, **match_row(row, odoo_by_model, sku_counts)})

    write_workbook(shopee_headers, output_rows, odoo_products, OUT_FILE)

    status_counts = Counter(row["mapping_status"] for row in output_rows)
    print(f"Wrote {OUT_FILE}")
    print(f"Shopee rows: {len(output_rows)}")
    print(f"Odoo rows: {len(odoo_products)}")
    for status, count in status_counts.most_common():
        print(f"{status}: {count}")
    print(f"Auto mapped: {sum(1 for row in output_rows if row['mapped_odoo_product_key'])}")


if __name__ == "__main__":
    main()
