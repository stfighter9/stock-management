#!/usr/bin/env python3
from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ACTUAL_STOCK = ROOT / "actual stock.xlsx"
DEFAULT_KEY_FILE = ROOT / "smart-sandbox-424309-n5-c615a9ae9610.json"

ALIAS_BY_DISPLAY_NAME = {
    "FO-031/PH SWEET CANDEE 0.6 Black tub 30s/1200": "FO-031/PH SWEET CANDEE 0.6 Black jar 30s/1200",
    "FO-FL01/PH 0.4 8 colors set/800": "FO-FL01/PH 0.4 8 colors set/100",
    "FO-GELB06 PAZTO 0.5 Blk Assorted box 12s/600": "FO-GELB06 PSLIDE PASTEL 0.5 Blk Assorted box 12s/600",
    "FO-GELE007/PH MAZZIC 0.5 Black box 12s/600": "FO-GELE007/PH FLEXMAZZIC 0.5 Black box 12s/600",
    "FO-GELE007/PH MAZZIC 0.5 Blue box 12s/600": "FO-GELE007/PH FLEXMAZZIC 0.5 Blue box 12s/600",
    "FO-PH01/XK SMART 0.7 Black box 10s/180": "FO-PH01/PH SMART 0.7 Black box 10s/180",
    "FO-PM05 Perma.Marker Black box 12s/600": "FO-PM05/XK Perma.Marker Black box 12s/600",
    "FO-WB025Whiteboard Marker HOSHI Black box 12s/600": "FO-WB025 Whiteboard Marker HOSHI Black box 12s/600",
}


def normalize(value) -> str:
    return "" if value is None else str(value).strip()


def build_service(key_file: Path):
    credentials = Credentials.from_service_account_file(str(key_file), scopes=["https://www.googleapis.com/auth/spreadsheets"])
    return build("sheets", "v4", credentials=credentials)


def read_products(service, spreadsheet_id: str) -> dict[str, str]:
    response = service.spreadsheets().values().get(spreadsheetId=spreadsheet_id, range="Products!A1:D").execute()
    values = response.get("values", [])
    if not values:
        return {}
    products = {}
    for row in values[1:]:
        if len(row) < 2:
            continue
        odoo_product_key = normalize(row[0])
        display_name = normalize(row[1])
        if display_name:
            products[display_name] = odoo_product_key
    return products


def ensure_empty_transactions(service, spreadsheet_id: str) -> None:
    response = service.spreadsheets().values().get(spreadsheetId=spreadsheet_id, range="Inventory_Transactions!A2:M").execute()
    values = response.get("values", [])
    non_empty = [row for row in values if any(normalize(cell) for cell in row)]
    if non_empty:
        raise RuntimeError("Inventory_Transactions already has data. Abort to avoid duplicate initial stock import.")


def build_transaction_rows(actual_stock_path: Path, products_by_display_name: dict[str, str], transaction_date: str, reference_id: str):
    workbook = load_workbook(actual_stock_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    header = [normalize(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    header_index = {name: idx for idx, name in enumerate(header)}

    rows = []
    unresolved = []
    aliases_used = []
    counter = 1

    for row in sheet.iter_rows(min_row=2, values_only=True):
        display_name = normalize(row[header_index["Display Name"]])
        if not display_name:
            continue
        actual_stock = row[header_index["actual stock"]]
        qty = 0 if actual_stock in (None, "") else float(actual_stock)
        if qty <= 0:
            continue

        resolved_display_name = ALIAS_BY_DISPLAY_NAME.get(display_name, display_name)
        odoo_product_key = products_by_display_name.get(resolved_display_name)
        if not odoo_product_key:
            unresolved.append((display_name, qty))
            continue

        if resolved_display_name != display_name:
            aliases_used.append((display_name, resolved_display_name, qty))

        transaction_id = f"INIT-STOCK-{transaction_date.replace('-', '')}-{counter:04d}"
        counter += 1
        qty_value = int(qty) if qty.is_integer() else qty
        rows.append(
            [
                transaction_id,
                transaction_date,
                "IN",
                odoo_product_key,
                qty_value,
                "STOCKTAKE",
                reference_id,
                "",
                "",
                "",
                "Không cần ghi nhận",
                "Codex",
                "Initial stock import from actual stock.xlsx",
            ]
        )

    return rows, aliases_used, unresolved


def append_transactions(service, spreadsheet_id: str, rows: list[list[object]]) -> None:
    if not rows:
        return
    service.spreadsheets().values().append(
        spreadsheetId=spreadsheet_id,
        range="Inventory_Transactions!A1",
        valueInputOption="USER_ENTERED",
        insertDataOption="INSERT_ROWS",
        body={"values": rows},
    ).execute()


def parse_args():
    parser = argparse.ArgumentParser(description="Import initial actual stock as IN transactions.")
    parser.add_argument("--spreadsheet-id", required=True)
    parser.add_argument("--actual-stock-file", default=str(DEFAULT_ACTUAL_STOCK))
    parser.add_argument("--key-file", default=str(DEFAULT_KEY_FILE))
    parser.add_argument("--transaction-date", default=str(date.today()))
    parser.add_argument("--reference-id", default=None)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    transaction_date = args.transaction_date
    reference_id = args.reference_id or f"INITIAL_ACTUAL_STOCK_{transaction_date}"

    service = build_service(Path(args.key_file).resolve())
    ensure_empty_transactions(service, args.spreadsheet_id)
    products_by_display_name = read_products(service, args.spreadsheet_id)
    rows, aliases_used, unresolved = build_transaction_rows(
        Path(args.actual_stock_file).resolve(),
        products_by_display_name,
        transaction_date,
        reference_id,
    )
    if unresolved:
        print("Unresolved rows:")
        for display_name, qty in unresolved:
            print(f"- {display_name} | qty={qty:g}")
        raise RuntimeError("Abort because some positive actual-stock rows could not be resolved to Products.")

    append_transactions(service, args.spreadsheet_id, rows)

    print(f"Appended transactions: {len(rows)}")
    print(f"Reference ID: {reference_id}")
    print(f"Aliases used: {len(aliases_used)}")
    for source_name, target_name, qty in aliases_used:
        print(f"ALIAS {source_name} -> {target_name} | qty={qty:g}")


if __name__ == "__main__":
    main()
